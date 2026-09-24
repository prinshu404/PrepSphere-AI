import re

from .models import Interview, LoginHistory
from django.contrib.auth import authenticate
from django.contrib.auth.models import User
from rest_framework.decorators import api_view
from rest_framework.response import Response

from pypdf import PdfReader
from docx import Document


def get_demo_user(request):
    authorization = request.headers.get("Authorization", "")

    if not authorization.startswith("Bearer "):
        return None

    token = authorization.replace("Bearer ", "", 1)

    prefix = "prepsphere-demo-token-"

    if not token.startswith(prefix):
        return None

    user_id = token.replace(prefix, "", 1)

    try:
        return User.objects.get(
            id=int(user_id),
            is_active=True,
        )
    except (User.DoesNotExist, ValueError):
        return None


@api_view(["GET"])
def home(request):
    total_students = User.objects.filter(is_active=True).count()

    return Response({
        "platform": {
            "name": "PrepSphere AI",
            "tagline": "Prepare smarter. Perform better.",
            "description": (
                "AI-powered interview and placement preparation "
                "for students."
            ),
        },
        "stats": {
            "students": total_students,
            "interviews": Interview.objects.count(),
            "questions": 0,
            "reports": 0,
        },
        "features": [
            {
                "title": "AI Mock Interviews",
                "description": (
                    "Practice interview questions and improve "
                    "your confidence with guided AI feedback."
                ),
            },
            {
                "title": "Resume Analyzer",
                "description": (
                    "Find gaps in your resume and get practical "
                    "suggestions to make it stronger."
                ),
            },
            {
                "title": "Coding Practice",
                "description": (
                    "Practice coding problems based on common "
                    "technical interview patterns."
                ),
            },
            {
                "title": "Performance Reports",
                "description": (
                    "Track your preparation and understand "
                    "where you need more practice."
                ),
            },
            {
                "title": "Career Roadmap",
                "description": (
                    "Follow a structured preparation path based "
                    "on your skills and career goals."
                ),
            },
            {
                "title": "Company Preparation",
                "description": (
                    "Prepare with company-focused topics and "
                    "common interview patterns."
                ),
            },
        ],
    })


@api_view(["POST"])
def login(request):
    email = request.data.get("email")
    password = request.data.get("password")

    if not email or not password:
        return Response({
            "message": "Email and password are required."
        }, status=400)

    user = authenticate(username=email, password=password)

    if user is None:
        LoginHistory.objects.create(
            email=email,
            success=False,
        )

        return Response({
            "message": "Invalid email or password."
        }, status=401)

    LoginHistory.objects.create(
        user=user,
        email=email,
        success=True,
    )

    return Response({
        "token": f"prepsphere-demo-token-{user.id}",
        "user": {
            "id": user.id,
            "email": user.email,
            "name": user.first_name,
        }
    })


@api_view(["POST"])
def register(request):
    name = request.data.get("name")
    email = request.data.get("email")
    password = request.data.get("password")

    if not name or not email or not password:
        return Response({
            "message": "Name, email and password are required."
        }, status=400)

    if User.objects.filter(username=email).exists():
        return Response({
            "message": "An account with this email already exists."
        }, status=400)

    user = User.objects.create_user(
        username=email,
        email=email,
        password=password,
        first_name=name,
    )

    return Response({
        "token": f"prepsphere-demo-token-{user.id}",
        "user": {
            "id": user.id,
            "email": user.email,
            "name": user.first_name,
        }
    }, status=201)


@api_view(["GET"])
def dashboard(request):
    total_students = User.objects.filter(is_active=True).count()
    total_logins = LoginHistory.objects.filter(success=True).count()
    failed_logins = LoginHistory.objects.filter(success=False).count()

    return Response({
        "message": "Welcome to PrepSphere AI!",
        "stats": {
            "students": total_students,
            "logins": total_logins,
            "failed_logins": failed_logins,
            "interviews": Interview.objects.count(),
            "completed": Interview.objects.filter(
                status="Completed"
            ).count(),
            "reports": 0,
        }
    })


@api_view(["GET", "POST"])
def interviews(request):
    user = get_demo_user(request)

    if user is None:
        return Response({
            "message": "Please log in again to continue."
        }, status=401)

    if request.method == "GET":
        interviews = Interview.objects.filter(
            user=user
        ).order_by("-created_at")

        data = []

        for interview in interviews:
            data.append({
                "id": interview.id,
                "title": interview.title,
                "interview_type": interview.interview_type,
                "status": interview.status,
                "score": interview.score,
                "created_at": interview.created_at,
            })

        return Response(data)

    title = request.data.get("title")
    interview_type = request.data.get(
        "interview_type",
        "Technical"
    )

    if not title:
        return Response({
            "message": "Interview title is required."
        }, status=400)

    interview = Interview.objects.create(
        user=user,
        title=title,
        interview_type=interview_type,
    )

    return Response({
        "id": interview.id,
        "title": interview.title,
        "interview_type": interview.interview_type,
        "status": interview.status,
        "score": interview.score,
        "created_at": interview.created_at,
    }, status=201)


def extract_pdf_text(file):
    reader = PdfReader(file)

    pages = []

    for page in reader.pages:
        text = page.extract_text() or ""

        if text.strip():
            pages.append(text)

    return "\n".join(pages)


def extract_docx_text(file):
    document = Document(file)

    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            cells = []

            for cell in row.cells:
                if cell.text.strip():
                    cells.append(cell.text.strip())

            if cells:
                paragraphs.append(" ".join(cells))

    return "\n".join(paragraphs)


def clean_resume_text(text):
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def analyze_resume_text(text):
    normalized_text = text.lower()

    skills_catalog = [
        "python",
        "java",
        "javascript",
        "typescript",
        "react",
        "angular",
        "vue",
        "django",
        "flask",
        "node.js",
        "express",
        "sql",
        "mysql",
        "postgresql",
        "mongodb",
        "html",
        "css",
        "git",
        "github",
        "docker",
        "kubernetes",
        "aws",
        "azure",
        "gcp",
        "rest api",
        "graphql",
        "machine learning",
        "data analysis",
        "pandas",
        "numpy",
        "tensorflow",
        "pytorch",
        "c++",
        "c#",
    ]

    skills = []

    for skill in skills_catalog:
        if skill in normalized_text:
            skills.append(skill)

    section_keywords = {
        "summary": [
            "summary",
            "profile",
            "objective",
        ],
        "experience": [
            "experience",
            "work experience",
            "professional experience",
            "employment",
        ],
        "education": [
            "education",
            "academic",
            "qualification",
        ],
        "skills": [
            "skills",
            "technical skills",
            "technologies",
        ],
        "projects": [
            "projects",
            "personal projects",
            "academic projects",
        ],
        "certifications": [
            "certification",
            "certifications",
        ],
    }

    detected_sections = []

    for section, keywords in section_keywords.items():
        if any(keyword in normalized_text for keyword in keywords):
            detected_sections.append(section)

    section_score = min(len(detected_sections) * 2, 10)

    email_found = bool(
        re.search(
            r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b",
            text,
            re.IGNORECASE,
        )
    )

    phone_found = bool(
        re.search(
            r"(?<!\d)(?:\+?\d[\d\s().-]{8,}\d)(?!\d)",
            text,
        )
    )

    contact_score = 5 if email_found and phone_found else 3 if (
        email_found or phone_found
    ) else 0

    achievement_patterns = [
        r"\b\d+%",
        r"\b\d+\+?\s*(?:users|projects|clients|students|members)\b",
        r"\b(?:increased|improved|reduced|saved|generated|achieved|grew)\b",
    ]

    achievement_matches = 0

    for pattern in achievement_patterns:
        achievement_matches += len(
            re.findall(
                pattern,
                normalized_text,
                re.IGNORECASE,
            )
        )

    achievement_score = min(achievement_matches * 2, 10)

    keyword_score = min(len(skills) * 1.5, 20)

    experience_score = 20 if "experience" in detected_sections else 0

    project_score = 15 if "projects" in detected_sections else 0

    formatting_score = 5

    total_score = round(
        keyword_score
        + section_score
        + contact_score
        + achievement_score
        + experience_score
        + project_score
        + formatting_score
    )

    total_score = max(0, min(total_score, 100))

    common_keywords = [
        "rest api",
        "testing",
        "docker",
        "git",
        "github",
        "sql",
        "cloud",
        "agile",
        "communication",
        "leadership",
    ]

    missing_keywords = [
        keyword
        for keyword in common_keywords
        if keyword not in normalized_text
    ]

    missing_keywords = missing_keywords[:5]

    suggestions = []

    if not skills:
        suggestions.append(
            "Add a dedicated technical skills section with relevant technologies."
        )
    elif len(skills) < 5:
        suggestions.append(
            "Add more relevant technical skills that match your target role."
        )

    if "experience" not in detected_sections:
        suggestions.append(
            "Add a clear experience section with role responsibilities and results."
        )

    if "projects" not in detected_sections:
        suggestions.append(
            "Add a projects section with technologies and measurable outcomes."
        )

    if achievement_matches == 0:
        suggestions.append(
            "Add measurable achievements using numbers, percentages, or concrete results."
        )

    if not email_found or not phone_found:
        suggestions.append(
            "Make sure your email address and phone number are clearly visible."
        )

    if missing_keywords:
        suggestions.append(
            "Consider adding relevant keywords such as "
            + ", ".join(missing_keywords[:3])
            + " when they accurately reflect your experience."
        )

    if not suggestions:
        suggestions.append(
            "Keep your resume focused, measurable, and aligned with the target role."
        )

    if total_score >= 85:
        summary = "Strong ATS compatibility with several relevant resume signals."
    elif total_score >= 70:
        summary = "Good ATS compatibility with some areas that can be improved."
    elif total_score >= 50:
        summary = "Moderate ATS compatibility with several areas needing improvement."
    else:
        summary = "The resume needs improvement in several ATS-related areas."

    return {
        "ats_score": total_score,
        "skills": skills,
        "missing_keywords": missing_keywords,
        "sections": detected_sections,
        "suggestions": suggestions,
        "summary": summary,
    }


@api_view(["POST"])
def resume_analyze(request):
    resume = request.FILES.get("resume")

    if resume is None:
        return Response({
            "message": "Please upload a resume file."
        }, status=400)

    allowed_extensions = (".pdf", ".doc", ".docx")
    file_name = resume.name.lower()

    if not file_name.endswith(allowed_extensions):
        return Response({
            "message": "Please upload a PDF, DOC, or DOCX resume."
        }, status=400)

    if resume.size > 5 * 1024 * 1024:
        return Response({
            "message": "Resume file must be 5 MB or smaller."
        }, status=400)

    try:
        if file_name.endswith(".pdf"):
            extracted_text = extract_pdf_text(resume)
        elif file_name.endswith(".docx"):
            extracted_text = extract_docx_text(resume)
        else:
            return Response({
                "message": (
                    "Legacy DOC files are not supported for text extraction yet. "
                    "Please upload a PDF or DOCX file."
                )
            }, status=400)

        extracted_text = clean_resume_text(extracted_text)

        if not extracted_text:
            return Response({
                "message": (
                    "No readable text was found in the resume. "
                    "Please upload a text-based PDF or DOCX file."
                )
            }, status=400)

        analysis = analyze_resume_text(extracted_text)

        return Response({
            "status": "analyzed",
            "file_name": resume.name,
            "file_size": resume.size,
            **analysis,
        })

    except Exception:
        return Response({
            "message": (
                "Unable to read this resume. "
                "Please upload a valid PDF or DOCX file."
            )
        }, status=400)